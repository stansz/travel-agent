#!/usr/bin/env python3
"""Travel Agent Chat — FastAPI backend with Pi agent, sessions, model/thinking controls."""

import asyncio
import os
import sqlite3
import urllib.parse
import uuid
from datetime import datetime

from fastapi import FastAPI, Request, Query
from fastapi.responses import HTMLResponse, JSONResponse, PlainTextResponse, Response
from docx import Document
from docx.shared import Pt, Inches
import io
import re

app = FastAPI(title="Travel Agent")

from starlette.middleware.base import BaseHTTPMiddleware

class NoCacheMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        response = await call_next(request)
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        return response

app.add_middleware(NoCacheMiddleware)

PI_CMD = "pi"
PI_PROVIDER = os.environ.get("PI_PROVIDER", "zai")

# glm-5.3-flash always engages in thinking: "off"/"minimal" are rejected by the
# provider (HTTP 400), so the defaults below must never be "off".
DEFAULT_MODEL = os.environ.get("PI_MODEL", "glm-5.3-flash")
PI_MODEL = DEFAULT_MODEL
DEFAULT_THINKING = os.environ.get("PI_THINKING", "high")
PI_SKILL = "/app/web/pi-skill.md"
HIMALAYA = os.environ.get("HIMALAYA", "himalaya")
DB_PATH = "/app/data/sessions.db"
LT_PORT = os.environ.get("LT_PORT", "5000")

os.makedirs("/app/data", exist_ok=True)

MODELS = ["glm-5.3-flash", "glm-5-turbo", "glm-5.1", "glm-4.7"]
THINKING_LEVELS = ["low", "medium", "high", "xhigh"]

with open(PI_SKILL) as f:
    TRAVEL_SYSTEM_PROMPT = f.read()


def get_db() -> sqlite3.Connection:
    db = sqlite3.connect(DB_PATH, timeout=10)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA journal_mode=WAL")
    db.execute("PRAGMA busy_timeout=5000")
    db.execute("PRAGMA foreign_keys=ON")
    return db


def init_db():
    db = get_db()
    db.executescript("""
        CREATE TABLE IF NOT EXISTS sessions (
            id TEXT PRIMARY KEY,
            title TEXT DEFAULT 'New Session',
            model TEXT DEFAULT 'glm-5.3-flash',
            thinking TEXT DEFAULT 'high',
            created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            session_id TEXT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
            role TEXT NOT NULL CHECK(role IN ('user','agent')),
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_messages_session ON messages(session_id);
    """)
    db.commit()
    db.close()


init_db()


# ── Session management endpoints ────────────────────────────────────────────

@app.get("/sessions")
async def list_sessions():
    db = get_db()
    rows = db.execute(
        "SELECT id, title, model, thinking, created_at FROM sessions ORDER BY created_at DESC"
    ).fetchall()
    db.close()
    return JSONResponse([dict(r) for r in rows])


@app.post("/sessions")
async def create_session():
    sid = uuid.uuid4().hex[:12]
    now = datetime.utcnow().isoformat()
    db = get_db()
    db.execute("INSERT INTO sessions(id, title, created_at) VALUES(?, ?, ?)", (sid, "New Session", now))
    db.commit()
    db.close()
    return JSONResponse({"id": sid, "title": "New Session", "model": DEFAULT_MODEL, "thinking": DEFAULT_THINKING, "created_at": now})


@app.get("/sessions/{sid}")
async def get_session(sid: str):
    db = get_db()
    session = db.execute("SELECT * FROM sessions WHERE id=?", (sid,)).fetchone()
    if not session:
        db.close()
        return JSONResponse({"error": "Not found"}, status_code=404)
    messages = db.execute(
        "SELECT role, content, created_at FROM messages WHERE session_id=? ORDER BY id", (sid,)
    ).fetchall()
    db.close()
    return JSONResponse({
        "session": dict(session),
        "messages": [dict(m) for m in messages]
    })


@app.delete("/sessions/{sid}")
async def delete_session(sid: str):
    db = get_db()
    db.execute("DELETE FROM messages WHERE session_id=?", (sid,))
    db.execute("DELETE FROM sessions WHERE id=?", (sid,))
    db.commit()
    db.close()
    return JSONResponse({"ok": True})


@app.patch("/sessions/{sid}")
async def update_session(sid: str, request: Request):
    body = await request.json()
    db = get_db()
    for field in ("model", "thinking", "title"):
        if field in body:
            db.execute(f"UPDATE sessions SET {field}=? WHERE id=?", (body[field], sid))
    db.commit()
    db.close()
    return JSONResponse({"ok": True})


# ── Chat / download / email ─────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse)
async def index():
    return HTMLResponse(HTML_PAGE)


@app.get("/test", response_class=HTMLResponse)
async def test_page():
    return HTMLResponse("""<!DOCTYPE html>
<html><head><meta charset="UTF-8"><title>Tab Test</title>
<style>
body{font-family:sans-serif;background:#0f172a;color:#e2e8f0;margin:0;padding:20px;}
.tab{background:transparent;border:1px solid #334155;color:#94a3b8;padding:8px 16px;cursor:pointer;border-radius:8px;margin-right:4px;}
.tab.active{background:#2563eb;color:#fff;border-color:#2563eb;}
.view{padding:20px;margin-top:12px;border-radius:8px;background:#1e293b;}
.hidden{display:none!important;}
</style></head><body>
<h2>Tab Test</h2>
<button class="tab active" id="tb-chat" onclick="sw('chat')">Chat</button>
<button class="tab" id="tb-translate" onclick="sw('translate')">Translate</button>
<div id="v-chat" class="view">✅ Chat view — you should see this by default</div>
<div id="v-translate" class="view hidden">✅ Translate view — this should appear when you click Translate</div>
<script>
function sw(t){
 document.querySelectorAll('.tab').forEach(function(x){x.classList.remove('active')});
 document.getElementById('tb-'+t).classList.add('active');
 document.getElementById('v-chat').classList.toggle('hidden',t!=='chat');
 document.getElementById('v-translate').classList.toggle('hidden',t!=='translate');
 console.log('switched to',t);
}
console.log('test page loaded');
</script>
</body></html>""")


@app.post("/chat")
async def chat(request: Request):
    body = await request.json()
    message = body.get("message", "").strip()
    session_id = body.get("session_id", "")
    model = body.get("model", PI_MODEL)
    thinking = body.get("thinking", DEFAULT_THINKING)

    if not message:
        return JSONResponse({"error": "empty message"}, status_code=400)

    if model not in MODELS:
        model = PI_MODEL
    if thinking not in THINKING_LEVELS:
        thinking = DEFAULT_THINKING

    # Auto-create session if needed
    db = get_db()
    if not session_id:
        sid = uuid.uuid4().hex[:12]
        title = message[:80]
        now = datetime.utcnow().isoformat()
        db.execute(
            "INSERT INTO sessions(id, title, model, thinking, created_at) VALUES(?,?,?,?,?)",
            (sid, title, model, thinking, now)
        )
        session_id = sid
    else:
        # Update session model/thinking
        db.execute("UPDATE sessions SET model=?, thinking=? WHERE id=?", (model, thinking, session_id))
        # Set title from first user message if still default
        row = db.execute("SELECT title FROM sessions WHERE id=?", (session_id,)).fetchone()
        if row and row["title"] == "New Session":
            db.execute("UPDATE sessions SET title=? WHERE id=?", (message[:80], session_id))

    # Save user message
    now = datetime.utcnow().isoformat()
    db.execute(
        "INSERT INTO messages(session_id, role, content, created_at) VALUES(?,?,?,?)",
        (session_id, "user", message, now)
    )
    db.commit()
    db.close()

    try:
        result = await run_pi(message, model, thinking)
    except Exception as e:
        result = f"Sorry, something went wrong: {e}"

    # Save agent response
    db = get_db()
    db.execute(
        "INSERT INTO messages(session_id, role, content, created_at) VALUES(?,?,?,?)",
        (session_id, "agent", result, datetime.utcnow().isoformat())
    )
    db.commit()
    db.close()

    return JSONResponse({"reply": result, "session_id": session_id})


@app.post("/email")
async def email(request: Request):
    body = await request.json()
    text = body.get("text", "").strip()
    to = body.get("to", "").strip()
    if not text or not to:
        return JSONResponse({"error": "Missing text or email address"}, status_code=400)

    try:
        await send_email(text, to)
        return JSONResponse({"ok": True})
    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)})


@app.get("/api/languages")
async def get_languages():
    """Proxy to LibreTranslate languages list."""
    import urllib.request
    import json
    try:
        req = urllib.request.Request(f"http://127.0.0.1:{LT_PORT}/languages")
        with urllib.request.urlopen(req, timeout=5) as resp:
            return JSONResponse(json.loads(resp.read()))
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=502)


@app.post("/api/translate")
async def translate_text(request: Request):
    """Proxy to LibreTranslate."""
    import urllib.request
    import json
    body = await request.json()
    text = body.get("text", "").strip()
    source = body.get("source", "auto")
    target = body.get("target", "en")
    if not text:
        return JSONResponse({"error": "No text provided"}, status_code=400)
    try:
        payload = json.dumps({"q": text, "source": source, "target": target, "format": "text"}).encode()
        req = urllib.request.Request(
            f"http://127.0.0.1:{LT_PORT}/translate",
            data=payload,
            headers={"Content-Type": "application/json"}
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            return JSONResponse(json.loads(resp.read()))
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=502)


@app.get("/download")
async def download(text: str = Query(""), fmt: str = Query("md")):
    if not text:
        return PlainTextResponse("No content", status_code=400)

    if fmt == "docx":
        docx_bytes = markdown_to_docx(text)
        headers = {"Content-Disposition": "attachment; filename=travel-info.docx"}
        return Response(docx_bytes, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers)
    else:
        headers = {"Content-Disposition": "attachment; filename=travel-info.md"}
        return PlainTextResponse(text, media_type="text/markdown", headers=headers)


def markdown_to_docx(md_text: str) -> bytes:
    """Convert Markdown to a simple DOCX with basic formatting."""
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    for line in md_text.split('\n'):
        if line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
        elif line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
        elif line.startswith('|') and '|' in line[1:]:
            if '---' in line.replace('|', '').replace('-', '').replace(' ', ''):
                continue
            cells = [c.strip() for c in line.split('|')[1:-1]]
            if cells:
                doc.add_paragraph('  \u2022  '.join(cells))
        elif line.strip().startswith(('\u2022 ', '- ', '* ')):
            doc.add_paragraph(line.strip(), style='List Bullet')
        elif line.strip().startswith(('1. ', '2. ', '3. ')):
            doc.add_paragraph(line.strip(), style='List Number')
        elif line.strip():
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                else:
                    p.add_run(part)
        else:
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def run_pi(message: str, model: str | None = None, thinking: str | None = None) -> str:
    model = model or PI_MODEL
    thinking = thinking or DEFAULT_THINKING
    cmd = [
        PI_CMD, "--provider", PI_PROVIDER, "--model", model,
        "--no-session", "--no-builtin-tools", "--no-skills",
        "--tools", "bash", "--system-prompt", PI_SKILL,
        "--thinking", thinking, "-p", message,
    ]
    proc = await asyncio.create_subprocess_exec(
        *cmd, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
    try:
        stdout, stderr = await asyncio.wait_for(proc.communicate(), timeout=120)
    except asyncio.TimeoutError:
        proc.kill()
        return "Sorry, the travel search took too long. Please try a more specific query."
    out = stdout.decode("utf-8", errors="replace").strip()
    err = stderr.decode("utf-8", errors="replace").strip()
    if not out:
        # Surface the real reason (provider/CLI errors go to stderr) instead of a
        # silent "no results".
        return f"⚠️ The agent returned nothing. Details: {err or 'no output (exit ' + str(proc.returncode) + ')'}"
    return out


async def send_email(text: str, to: str) -> None:
    """Send markdown email via himalaya CLI."""
    subject = "Travel Info"
    first_line = text.strip().split("\n")[0]
    if first_line.startswith("#"):
        subject = first_line.lstrip("#").strip()[:80]
    elif first_line.startswith("**") and "**" in first_line[2:]:
        subject = first_line.strip("*").strip()[:80]

    proc = await asyncio.create_subprocess_exec(
        HIMALAYA, "write",
        "--subject", subject,
        "--header", "Content-Type: text/markdown; charset=utf-8",
        to,
        stdin=asyncio.subprocess.PIPE,
        stdout=asyncio.subprocess.PIPE,
        stderr=asyncio.subprocess.PIPE,
    )
    stdout, stderr = await asyncio.wait_for(
        proc.communicate(input=text.encode("utf-8")), timeout=30)
    if proc.returncode != 0:
        err = stderr.decode("utf-8", errors="replace").strip() or "Unknown error"
        raise RuntimeError(f"Himalaya failed: {err}")


HTML_PAGE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Travel Agent ✈️</title>
<script src="https://cdn.jsdelivr.net/npm/marked/marked.min.js"></script>
<style>
* { box-sizing: border-box; margin: 0; padding: 0; }
body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #0f172a; color: #e2e8f0; height: 100vh; overflow: hidden; }
#app { display: flex; height: 100vh; }
/* ── Sidebar ────────────────────────────── */
#sidebar { width: 220px; min-width: 220px; background: #1a1f2e; border-right: 1px solid #334155; display: flex; flex-direction: column; transition: margin-left 0.25s ease, opacity 0.25s ease; }
#sidebar.collapsed { margin-left: -220px; opacity: 0; pointer-events: none; }
.sidebar-header { padding: 14px 16px; border-bottom: 1px solid #1e293b; display: flex; align-items: center; justify-content: space-between; }
.sidebar-header h2 { font-size: 13px; font-weight: 600; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.5px; }
.sidebar-header button { background: none; border: 1px solid #334155; color: #64748b; border-radius: 5px; padding: 3px 7px; font-size: 14px; cursor: pointer; }
.sidebar-header button:hover { border-color: #2563eb; color: #e2e8f0; background: #1e293b; }
#sessionList { flex: 1; overflow-y: auto; padding: 8px; display: flex; flex-direction: column; gap: 2px; }
.session-item { padding: 10px 12px; border-radius: 8px; cursor: pointer; display: flex; align-items: center; gap: 8px; font-size: 13px; color: #cbd5e1; transition: background 0.12s; border: 1px solid transparent; }
.session-item:hover { background: #1e293b; }
.session-item.active { background: #1e293b; border-color: #2563eb; }
.session-item .title { flex: 1; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
.session-item .del { color: #64748b; font-size: 11px; padding: 2px 5px; border-radius: 4px; cursor: pointer; display: none; }
.session-item:hover .del { display: block; }
.session-item .del:hover { color: #ef4444; background: rgba(239,68,68,0.15); }
/* ── Main ────────────────────────────────── */
#main { flex: 1; display: flex; flex-direction: column; min-width: 0; }
header { background: #1e293b; padding: 10px 20px; border-bottom: 1px solid #334155; display: flex; align-items: center; gap: 12px; }
header h1 { font-size: 17px; font-weight: 600; }
header span.icon { font-size: 20px; }
header .hamburger { background: none; border: 1px solid #334155; color: #64748b; border-radius: 6px; padding: 5px 9px; font-size: 16px; cursor: pointer; margin-right: 4px; }
header .hamburger:hover { border-color: #2563eb; color: #e2e8f0; }
header .tabs { display: flex; gap: 4px; margin-left: 16px; }
header .tab { padding: 6px 14px; border-radius: 8px; border: 1px solid transparent; background: transparent; color: #64748b; font-size: 13px; cursor: pointer; transition: all 0.15s; }
header .tab:hover { color: #e2e8f0; background: #0f172a; }
header .tab.active { color: #e2e8f0; background: #0f172a; border-color: #334155; }
.lang-btn { background: none; border: 1px solid #334155; border-radius: 6px; color: #64748b; padding: 4px 10px; font-size: 12px; font-weight: 600; cursor: pointer; margin-left: auto; font-family: inherit; }
.lang-btn:hover { border-color: #2563eb; color: #e2e8f0; }
#chat-view, #translate-view { flex: 1; overflow-y: auto; display: flex; flex-direction: column; }
#translate-view { padding: 24px; gap: 16px; align-items: center; }
.hidden-view { display: none !important; }
.translate-box { width: 100%; max-width: 600px; display: flex; flex-direction: column; gap: 12px; }
.translate-box .lang-row { display: flex; align-items: center; gap: 10px; }
.translate-box select { padding: 7px 12px; border-radius: 6px; border: 1px solid #334155; background: #0f172a; color: #e2e8f0; font-size: 13px; cursor: pointer; }
.translate-box select:focus { outline: none; border-color: #2563eb; }
.translate-box .arrow { color: #64748b; font-size: 16px; }
.translate-box textarea { width: 100%; height: 120px; padding: 12px; border-radius: 8px; border: 1px solid #334155; background: #0f172a; color: #e2e8f0; font-size: 15px; resize: vertical; font-family: inherit; outline: none; }
.translate-box textarea:focus { border-color: #2563eb; }
.translate-box .tr-btn { padding: 10px 24px; border-radius: 8px; border: none; background: #2563eb; color: #fff; font-size: 14px; font-weight: 600; cursor: pointer; align-self: flex-end; }
.translate-box .tr-btn:hover { background: #1d4ed8; }
.translate-box .tr-btn:disabled { opacity: 0.5; cursor: not-allowed; }
.translate-result { width: 100%; max-width: 600px; padding: 16px; border-radius: 8px; background: #1e293b; border: 1px solid #334155; font-size: 16px; line-height: 1.6; display: none; }
.translate-result.show { display: block; animation: fadeIn 0.2s; }
.translate-error { color: #ef4444; font-size: 13px; }
#chat { flex: 1; overflow-y: auto; padding: 20px 24px; display: flex; flex-direction: column; gap: 16px; }
.msg { max-width: 75%; padding: 12px 16px; border-radius: 12px; line-height: 1.5; position: relative; animation: fadeIn 0.2s; }
@keyframes fadeIn { from { opacity: 0; transform: translateY(6px); } to { opacity: 1; transform: translateY(0); } }
.msg.user { align-self: flex-end; background: #2563eb; color: #fff; }
.msg.agent { align-self: flex-start; background: #1e293b; border: 1px solid #334155; }
.msg.agent h1, .msg.agent h2, .msg.agent h3 { margin: 8px 0 4px; font-size: 1em; }
.msg.agent ul, .msg.agent ol { padding-left: 20px; margin: 4px 0; }
.msg.agent li { margin: 2px 0; }
.msg.agent code { background: #0f172a; padding: 1px 6px; border-radius: 4px; font-size: 0.9em; font-family: 'SF Mono', Monaco, monospace; }
.msg.agent pre { background: #0f172a; padding: 8px 12px; border-radius: 8px; overflow-x: auto; margin: 4px 0; font-size: 0.85em; }
.msg.agent table { width: 100%; border-collapse: collapse; margin: 8px 0; font-size: 0.85em; }
.msg.agent th, .msg.agent td { border: 1px solid #334155; padding: 6px 10px; text-align: left; }
.msg.agent th { background: #1a2035; font-weight: 600; }
.msg.agent a { color: #60a5fa; }
.msg.agent blockquote { border-left: 3px solid #334155; margin: 4px 0; padding-left: 12px; color: #94a3b8; }
.msg.agent img { max-width: 100%; border-radius: 8px; }
.msg.agent hr { border: none; border-top: 1px solid #334155; margin: 8px 0; }
.msg.loading { align-self: flex-start; color: #64748b; font-style: italic; }
.msg-actions { margin-top: 10px; display: flex; gap: 8px; opacity: 0; transition: opacity 0.2s; }
.msg:hover .msg-actions, .msg-actions.show { opacity: 1; }
.msg-btn { padding: 4px 10px; border-radius: 6px; border: 1px solid #334155; background: transparent; color: #64748b; font-size: 12px; cursor: pointer; transition: all 0.15s; }
.msg-btn:hover { background: #1e293b; border-color: #2563eb; color: #e2e8f0; }
.msg-btn:disabled { opacity: 0.3; cursor: not-allowed; }
.email-popup { display: none; position: fixed; top: 0; left: 0; right: 0; bottom: 0; background: rgba(0,0,0,0.6); z-index: 100; align-items: center; justify-content: center; }
.email-popup.open { display: flex; }
.email-popup-inner { background: #1e293b; padding: 20px; border-radius: 12px; border: 1px solid #334155; width: 340px; }
.email-popup-inner input { width: 100%; padding: 10px; border-radius: 8px; border: 1px solid #334155; background: #0f172a; color: #e2e8f0; font-size: 14px; margin: 8px 0; }
.email-popup-inner button { padding: 8px 16px; border-radius: 8px; border: none; font-size: 13px; cursor: pointer; margin-right: 8px; }
.email-popup-inner .send-btn { background: #2563eb; color: #fff; }
.email-popup-inner .cancel-btn { background: #334155; color: #94a3b8; }
.preset-buttons { padding: 8px 20px; display: flex; gap: 8px; flex-wrap: wrap; }
.preset-buttons button { padding: 5px 12px; border-radius: 16px; border: 1px solid #334155; background: transparent; color: #94a3b8; font-size: 12px; cursor: pointer; white-space: nowrap; }
.preset-buttons button:hover { background: #1e293b; border-color: #2563eb; color: #e2e8f0; }
#input-bar { padding: 10px 20px; background: #1e293b; border-top: 1px solid #334155; display: flex; align-items: flex-end; gap: 10px; flex-wrap: wrap; }
.input-controls { display: flex; align-items: center; gap: 8px; }
@media (max-width: 600px) {
  /* Sidebar: overlay with backdrop */
  #sidebar { position: fixed; top: 0; left: 0; bottom: 0; z-index: 100; width: 240px; padding-top: 44px; }
  #sidebar.collapsed { margin-left: -240px; }
  .sidebar-backdrop { display: none; position: fixed; inset: 0; background: rgba(0,0,0,0.5); z-index: 99; }
  .sidebar-backdrop.show { display: block; }
  header .hamburger { position: relative; z-index: 101; }
  /* Header: compact */
  header { padding: 8px 12px; gap: 6px; }
  header h1 { font-size: 15px; }
  header .hamburger { font-size: 14px; padding: 4px 7px; }
  header .tabs { margin-left: auto; }
  header .tab { padding: 5px 10px; font-size: 12px; }
  /* Chat area: less padding */
  #chat { padding: 12px; }
  .msg { max-width: 90%; font-size: 13px; }
  /* Input bar: prominent, two-row */
  #input-bar { gap: 10px; padding: 10px 12px; flex-direction: column; align-items: stretch; }
  .input-controls { width: 100%; justify-content: space-between; }
  .input-controls select { flex: 1; max-width: 50%; font-size: 13px; padding: 7px 10px; }
  .thinking-seg button { font-size: 12px; padding: 6px 10px; }
  #input-bar textarea { flex: none; width: 100%; font-size: 16px; padding: 12px 14px; min-height: 200px; }
  #input-bar button.send-btn { align-self: flex-end; padding: 10px 20px; font-size: 15px; }
  /* Preset buttons: compact */
  .preset-buttons { padding: 6px 10px; gap: 5px; }
  .preset-buttons button { padding: 4px 9px; font-size: 11px; }
}
#input-bar textarea { flex: 1; padding: 10px 14px; border-radius: 8px; border: 1px solid #334155; background: #0f172a; color: #e2e8f0; font-size: 14px; outline: none; min-width: 0; min-height: 300px; max-height: 400px; resize: none; font-family: inherit; line-height: 1.4; overflow-y: auto; }
#input-bar textarea:focus { border-color: #2563eb; }
#input-bar textarea::placeholder { color: #64748b; }
#input-bar button.send-btn { padding: 9px 18px; border-radius: 8px; border: none; background: #2563eb; color: #fff; font-size: 14px; font-weight: 600; cursor: pointer; white-space: nowrap; }
#input-bar button.send-btn:hover { background: #1d4ed8; }
#input-bar button.send-btn:disabled { opacity: 0.5; cursor: not-allowed; }
#input-bar select { padding: 6px 10px; border-radius: 6px; border: 1px solid #334155; background: #0f172a; color: #94a3b8; font-size: 12px; cursor: pointer; }
#input-bar select:focus { outline: none; border-color: #2563eb; }
.thinking-seg { display: flex; border-radius: 6px; overflow: hidden; border: 1px solid #334155; }
.thinking-seg button { border: none; border-radius: 0; padding: 4px 8px; font-size: 11px; background: #0f172a; color: #64748b; cursor: pointer; }
.thinking-seg button:not(:last-child) { border-right: 1px solid #334155; }
.thinking-seg button.active { background: #2563eb; color: #fff; }
.toast { position: fixed; bottom: 20px; left: 50%; transform: translateX(-50%); background: #2563eb; color: #fff; padding: 10px 20px; border-radius: 8px; font-size: 14px; z-index: 200; opacity: 0; transition: opacity 0.3s; pointer-events: none; }
.toast.show { opacity: 1; }
</style>
</head>
<body>
<div id="app">
  <!-- ── Sidebar ─────────────────────────── -->
  <div id="sidebar">
    <div class="sidebar-header">
      <h2>Sessions</h2>
      <button onclick="newSession()" title="New session">+</button>
    </div>
    <div id="sessionList"></div>
  </div>

  <!-- ── Main ─────────────────────────────── -->
  <div id="main">
    <header>
      <button class="hamburger" onclick="toggleSidebar()" title="Toggle sidebar">☰</button>
      <span class="icon">✈️</span>
      <h1>Travel Agent</h1>
      <div class="tabs">
        <button class="tab active" id="tab-chat" onclick="switchTab('chat')">💬 Chat</button>
        <button class="tab" id="tab-translate" onclick="switchTab('translate')">🌐 Translate</button>
      </div>
      <button class="lang-btn" id="langToggle" onclick="toggleLang()" title="Switch language">EN</button>
    </header>

    <div id="chat-view">
    <div id="chat">
      <div class="msg agent">
        👋 Hi! I'm your AI travel assistant. I can help with:<br><br>
        ✈️ <b>Flight prices &amp; status</b> from Vancouver<br>
        🏨 <b>Hotel search</b> at your destination<br>
        🌤️ <b>Weather forecasts</b><br>
        ⚠️ <b>Travel advisories</b> (Canadian government)<br>
        💱 <b>Currency conversion</b><br>
        🔌 <b>Power plugs &amp; emergency numbers</b><br>
        🗣️ <b>Translation &amp; travel phrases</b><br><br>
        What can I help you plan?
      </div>
    </div>

    <div class="preset-buttons" id="presets">
      <button onclick="ask('Show me YVR flight deals')">🎫 YVR Deals</button>
      <button id="dyn1" onclick="ask('Flights from YVR to Tokyo')">🇯🇵 Tokyo Flights</button>
      <button id="dyn2" onclick="ask('Weather in Bangkok this week')">🌤️ Bangkok</button>
      <button id="dyn3" onclick="ask('Travel advisory for Thailand')">⚠️ Thailand</button>
      <button id="dyn4" onclick="ask('What to see in Japan? WikiVoyage')">📖 Japan Guide</button>
      <button onclick="ask('Recommend me a trip for this summer, budget friendly')">🧳 Recommend a Trip</button>
      <button onclick="ask('What are some fun unique trips from Vancouver? Search Reddit')">🔍 Reddit Trips</button>
      <button onclick="ask('Show me cruise deals from Canada')">🚢 Cruise Deals</button>
    </div>

    <div id="input-bar">
      <textarea id="input" rows="1" placeholder="Ask about flights, weather, advisories..." onkeydown="if(event.key=='Enter'&&!event.shiftKey){event.preventDefault();send()}" oninput="this.style.height='auto';this.style.height=Math.min(this.scrollHeight,400)+'px'" autofocus></textarea>
      <button class="send-btn" id="send-btn" onclick="send()">Send</button>
    </div>
    </div>

    <div id="translate-view" class="hidden-view">
      <div class="translate-box">
        <div class="lang-row">
          <select id="srcLang"><option value="auto">Auto-detect</option></select>
          <span class="arrow">→</span>
          <select id="tgtLang"></select>
          <button class="tr-btn" onclick="doTranslate()">Translate</button>
        </div>
        <textarea id="trInput" placeholder="Enter text to translate..."></textarea>
        <div class="translate-result" id="trResult"></div>
        <div class="translate-error" id="trError"></div>
      </div>
    </div>
  </div>
</div>

<!-- Popups -->
<div class="email-popup" id="emailPopup">
  <div class="email-popup-inner">
    <div style="color:#e2e8f0;font-weight:600;margin-bottom:4px;">✉️ Email this info</div>
    <input type="email" id="emailAddr" placeholder="you@example.com">
    <div style="margin-top: 12px;">
      <button class="send-btn" onclick="doEmail()">Send</button>
      <button class="cancel-btn" onclick="closeEmail()">Cancel</button>
    </div>
  </div>
</div>
<div class="sidebar-backdrop" id="sidebarBackdrop" onclick="toggleSidebar()"></div>
<div class="toast" id="toast"></div>

<script>
const chat = document.getElementById('chat');
const input = document.getElementById('input');
const btn = document.getElementById('send-btn');
const sessionList = document.getElementById('sessionList');
const langBtn = document.getElementById('langToggle');
let currentSessionId = '';
let emailTargetText = '';
let currentLang = 'en';

// ── Sidebar toggle ────────────────────────────────────────────────────────
function toggleSidebar() {
  const sb = document.getElementById('sidebar');
  const bd = document.getElementById('sidebarBackdrop');
  sb.classList.toggle('collapsed');
  if (bd) bd.classList.toggle('show', !sb.classList.contains('collapsed'));
}

// ── Language / i18n ─────────────────────────────────────────────────────────

const LANG = {
  en: {
    title: 'Travel Agent',
    tabChat: '💬 Chat',
    tabTrans: '🌐 Translate',
    inputPlaceholder: 'Ask about flights, weather, advisories...',
    sendBtn: 'Send',
    sidebarTitle: 'Sessions',
    newSession: '🆕 New session started. What can I help you plan?',
    deleteConfirm: 'Delete this session?',
    deleted: '🗑 Session deleted. Start a new one!',
    loading: 'Searching...',
    noResponse: 'No response',
    emailTitle: '✉️ Email this info',
    emailPlaceholder: 'you@example.com',
    emailSend: 'Send',
    emailCancel: 'Cancel',
    emailSent: 'Email sent! ✉️',
    trBtn: 'Translate',
    trPlaceholder: 'Enter text to translate...',
    trAuto: 'Auto-detect',
    chatWelcome: `👋 Hi! I'm your AI travel assistant. I can help with:<br><br>
✈️ <b>Flight prices &amp; status</b> from Vancouver<br>
🏨 <b>Hotel search</b> at your destination<br>
🌤️ <b>Weather forecasts</b><br>
⚠️ <b>Travel advisories</b> (Canadian government)<br>
💱 <b>Currency conversion</b><br>
🔌 <b>Power plugs &amp; emergency numbers</b><br>
🗣️ <b>Translation &amp; travel phrases</b><br><br>
What can I help you plan?`,
  },
  zh: {
    title: '旅行助手',
    tabChat: '💬 聊天',
    tabTrans: '🌐 翻译',
    inputPlaceholder: '询问航班、天气、旅行建议...',
    sendBtn: '发送',
    sidebarTitle: '历史会话',
    newSession: '🆕 新会话已开始。有什么可以帮您规划？',
    deleteConfirm: '删除这个会话？',
    deleted: '🗑 会话已删除。开始一个新会话！',
    loading: '搜索中...',
    noResponse: '无回复',
    emailTitle: '✉️ 邮件发送',
    emailPlaceholder: 'you@example.com',
    emailSend: '发送',
    emailCancel: '取消',
    emailSent: '邮件已发送！✉️',
    trBtn: '翻译',
    trPlaceholder: '输入要翻译的文字...',
    trAuto: '自动检测',
    chatWelcome: `👋 您好！我是您的AI旅行助手。我可以帮助您：<br><br>
✈️ <b>查询航班价格和状态</b>（温哥华出发）<br>
🏨 <b>目的地酒店搜索</b><br>
🌤️ <b>天气预报</b><br>
⚠️ <b>旅行建议</b>（加拿大政府）<br>
💱 <b>货币换算</b><br>
🔌 <b>电源插头和紧急电话</b><br>
🗣️ <b>翻译和旅行常用语</b><br><br>
有什么可以帮您？`,
  },
};

function t(key) { return LANG[currentLang][key] || key; }

function toggleLang() {
  currentLang = currentLang === 'en' ? 'zh' : 'en';
  applyLang();
}

function applyLang() {
  langBtn.textContent = currentLang === 'en' ? 'EN' : '中文';
  document.title = t('title');
  document.querySelector('header h1').textContent = '✈ ' + t('title');
  document.getElementById('tab-chat').textContent = t('tabChat');
  document.getElementById('tab-translate').textContent = t('tabTrans');
  input.placeholder = t('inputPlaceholder');
  document.getElementById('send-btn').textContent = t('sendBtn');
  document.querySelector('.sidebar-header h2').textContent = t('sidebarTitle');
  // Translate tab
  const trBtn = document.querySelector('.tr-btn');
  if (trBtn) trBtn.textContent = t('trBtn');
  const trInput = document.getElementById('trInput');
  if (trInput) trInput.placeholder = t('trPlaceholder');
  const srcLang = document.getElementById('srcLang');
  if (srcLang) srcLang.options[0].text = t('trAuto');
  // Email popup
  const epTitle = document.querySelector('.email-popup-inner div');
  if (epTitle) epTitle.innerHTML = t('emailTitle');
  document.getElementById('emailAddr').placeholder = t('emailPlaceholder');
  // Re-render welcome message and presets
  const chatEl = document.getElementById('chat');
  chatEl.innerHTML = `<div class="msg agent">${t('chatWelcome')}</div>`;
  shufflePrompts();
  // Email popup buttons
  const epBtns = document.querySelectorAll('.email-popup-inner button');
  epBtns.forEach(b => {
    if (b.classList.contains('send-btn')) b.textContent = t('emailSend');
    if (b.classList.contains('cancel-btn')) b.textContent = t('emailCancel');
  });
}

// ── Session management ─────────────────────────────────────────────────────

async function refreshSessions() {
  const resp = await fetch('/sessions');
  const sessions = await resp.json();
  sessionList.innerHTML = '';
  sessions.forEach(s => {
    const div = document.createElement('div');
    div.className = 'session-item' + (s.id === currentSessionId ? ' active' : '');
    div.onclick = () => loadSessionById(s.id);
    div.innerHTML = `<span class="title">${escapeHtml(s.title.substring(0, 30))}</span>
      <span class="del" onclick="event.stopPropagation();deleteSession('${s.id}')" title="Delete">×</span>`;
    sessionList.appendChild(div);
  });
  // Auto-select first session only if none selected
  if (sessions.length > 0 && !currentSessionId) {
    const first = sessions[0];
    currentSessionId = first.id;
    sessionList.querySelector('.session-item').classList.add('active');
  }
}

function escapeHtml(s) {
  const d = document.createElement('div');
  d.textContent = s;
  return d.innerHTML;
}

async function newSession() {
  const resp = await fetch('/sessions', {method:'POST'});
  const s = await resp.json();
  currentSessionId = s.id;
  chat.innerHTML = `<div class="msg agent">${t('newSession')}</div>`;
  await refreshSessions();
}

async function loadSessionById(sid) {
  if (sid === currentSessionId) return;
  currentSessionId = sid;
  const resp = await fetch('/sessions/' + sid);
  const data = await resp.json();
  if (data.error) return;
  chat.innerHTML = '';
  data.messages.forEach(m => addMsg(m.content, m.role));
  chat.scrollTop = chat.scrollHeight;
  // Update active class
  sessionList.querySelectorAll('.session-item').forEach(el => {
    el.classList.toggle('active', el.querySelector('.title').textContent === data.session.title.substring(0, 30));
  });
  // Auto-close sidebar on mobile
  if (window.innerWidth <= 600) toggleSidebar();
}

async function deleteSession(sid) {
  if (!confirm(t('deleteConfirm'))) return;
  await fetch('/sessions/' + sid, {method:'DELETE'});
  if (sid === currentSessionId) {
    currentSessionId = '';
    chat.innerHTML = `<div class="msg agent">${t('deleted')}</div>`;
  }
  await refreshSessions();
}

// ── Chat ────────────────────────────────────────────────────────────────────

function addMsg(text, role) {
  const div = document.createElement('div');
  div.className = 'msg ' + role;
  if (role === 'agent') {
    div.innerHTML = marked.parse(text);
  } else {
    div.textContent = text;
  }

  if (role === 'agent' && text.length > 10) {
    const rawText = text;
    const actions = document.createElement('div');
    actions.className = 'msg-actions';
    actions.innerHTML = '<button class="msg-btn" title="Download as Markdown" onclick="downloadMd(this)">📥 .md</button> <button class="msg-btn" title="Download as Word Doc" onclick="downloadDocx(this)">📄 .docx</button> <button class="msg-btn" title="Email this" onclick="promptEmail(this)">✉️ Email</button>';
    actions.dataset.text = rawText;
    div.appendChild(actions);
  }

  chat.appendChild(div);
  chat.scrollTop = chat.scrollHeight;
  return div;
}

async function send() {
  const msg = input.value.trim();
  if (!msg) return;
  input.value = '';
  input.style.height = 'auto';
  btn.disabled = true;
  addMsg(msg, 'user');
  const loading = addMsg(t('loading'), 'loading');
  // Prepend language instruction for Chinese mode
  const prompt = currentLang === 'zh' ? '(请用中文回复) ' + msg : msg;
  try {
    const resp = await fetch('/chat', {
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({
        message: prompt,
        session_id: currentSessionId,
        model: '__DEFAULT_MODEL__',
        thinking: '__DEFAULT_THINKING__'
      })
    });
    const data = await resp.json();
    loading.remove();
    if (data.session_id) {
      currentSessionId = data.session_id;
      await refreshSessions();
    }
    addMsg(data.reply || t('noResponse'), 'agent');
  } catch(e) { loading.remove(); addMsg('Error: '+e.message, 'agent'); }
  btn.disabled = false;
  input.focus();
}

function downloadMd(btn) {
  const text = btn.parentElement.dataset.text;
  const blob = new Blob([text], {type:'text/markdown'});
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a');
  a.href = url; a.download = 'travel-info.md'; a.click();
  URL.revokeObjectURL(url);
}

async function downloadDocx(btn) {
  const text = btn.parentElement.dataset.text;
  const resp = await fetch('/download?fmt=docx&text=' + encodeURIComponent(text));
  const blob = await resp.blob();
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a');
  a.href = url; a.download = 'travel-info.docx'; a.click();
  URL.revokeObjectURL(url);
}

function promptEmail(btn) {
  emailTargetText = btn.parentElement.dataset.text;
  document.getElementById('emailPopup').classList.add('open');
  document.getElementById('emailAddr').focus();
}

function closeEmail() { document.getElementById('emailPopup').classList.remove('open'); }

async function doEmail() {
  const addr = document.getElementById('emailAddr').value.trim();
  if (!addr) return;
  closeEmail();
  try {
    const resp = await fetch('/email', {method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({text:emailTargetText,to:addr})});
    const data = await resp.json();
    toast(data.ok ? t('emailSent') : 'Failed: ' + (data.error||'unknown'));
  } catch(e) { toast('Error: '+e.message); }
}

function toast(msg) {
  const t = document.getElementById('toast');
  t.textContent = msg; t.classList.add('show');
  setTimeout(() => t.classList.remove('show'), 3000);
}

function ask(q) { input.value = q; send(); }

// ── Translate Tab ──────────────────────────────────────────────────────────

let languagesLoaded = false;

async function loadLanguages() {
  if (languagesLoaded) return;
  try {
    const resp = await fetch('/api/languages');
    const langs = await resp.json();
    if (langs.error) return;
    // Sort by name
    langs.sort((a, b) => a.name.localeCompare(b.name));
    const src = document.getElementById('srcLang');
    const tgt = document.getElementById('tgtLang');
    langs.forEach(l => {
      const opt = document.createElement('option');
      opt.value = l.code;
      opt.textContent = l.name;
      tgt.appendChild(opt.cloneNode(true));
      // Don't add to source (keep auto-detect only)
    });
    // Default target
    const enOpt = tgt.querySelector('option[value="en"]');
    if (enOpt) enOpt.selected = true;
    languagesLoaded = true;
  } catch(e) { console.error('Failed to load languages:', e); }
}

function switchTab(tab) {
  document.querySelectorAll('.tab').forEach(function(t){ t.classList.remove('active'); });
  document.getElementById('tab-' + tab).classList.add('active');
  document.getElementById('chat-view').classList.toggle('hidden-view', tab !== 'chat');
  document.getElementById('translate-view').classList.toggle('hidden-view', tab !== 'translate');
  if (tab === 'translate') {
    loadLanguages();
    document.getElementById('trInput').focus();
  } else {
    document.getElementById('input').focus();
  }
}

async function doTranslate() {
  const text = document.getElementById('trInput').value.trim();
  if (!text) return;
  const source = document.getElementById('srcLang').value;
  const target = document.getElementById('tgtLang').value;
  const resultEl = document.getElementById('trResult');
  const errorEl = document.getElementById('trError');
  resultEl.classList.remove('show');
  errorEl.textContent = '';

  try {
    const resp = await fetch('/api/translate', {
      method: 'POST',
      headers: {'Content-Type': 'application/json'},
      body: JSON.stringify({text, source, target})
    });
    const data = await resp.json();
    if (data.error) {
      errorEl.textContent = 'Error: ' + data.error;
    } else {
      resultEl.textContent = data.translatedText;
      resultEl.classList.add('show');
    }
  } catch(e) {
    errorEl.textContent = 'Error: ' + e.message;
  }
}

// ── Init ────────────────────────────────────────────────────────────────────

// Dynamic preset buttons — rotate destinations
const DYN_PROMPTS = [
  { emoji: '🌤️', labelEn: 'Bangkok Weather', askEn: 'Weather in Bangkok this week', labelZh: '曼谷天气', askZh: '曼谷这周天气怎么样' },
  { emoji: '⚠️', labelEn: 'Thailand Advisory', askEn: 'Travel advisory for Thailand', labelZh: '泰国旅行建议', askZh: '泰国的旅行安全建议' },
  { emoji: '🇯🇵', labelEn: 'Tokyo Flights', askEn: 'Flights from YVR to Tokyo under $1000', labelZh: '东京机票', askZh: '从温哥华到东京的航班价格' },
  { emoji: '📖', labelEn: 'Japan Guide', askEn: 'Tell me about visiting Japan - sights, food, culture', labelZh: '日本旅游攻略', askZh: '介绍日本旅游景点、美食和文化' },
  { emoji: '🇲🇽', labelEn: 'Mexico City Flights', askEn: 'Flights from YVR to Mexico City in July', labelZh: '墨西哥城机票', askZh: '温哥华到墨西哥城的七月机票' },
  { emoji: '🇰🇷', labelEn: 'Seoul Flights', askEn: 'Flights from YVR to Seoul under $1200', labelZh: '首尔机票', askZh: '温哥华到首尔的机票价格' },
  { emoji: '🇬🇧', labelEn: 'London Guide', askEn: 'Tell me about visiting London - what to see and do', labelZh: '伦敦旅游攻略', askZh: '伦敦有哪些必看景点和活动' },
  { emoji: '🇫🇷', labelEn: 'Paris Weather', askEn: 'Weather in Paris this week', labelZh: '巴黎天气', askZh: '巴黎这周的天气' },
  { emoji: '🇮🇹', labelEn: 'Italy Guide', askEn: 'What should I know about visiting Italy? WikiVoyage guide', labelZh: '意大利旅游攻略', askZh: '去意大利旅游需要知道什么' },
  { emoji: '🇻🇳', labelEn: 'Vietnam Flights', askEn: 'Flights from YVR to Vietnam this fall', labelZh: '越南机票', askZh: '温哥华到越南的秋季机票' },
  { emoji: '🇹🇭', labelEn: 'Phuket Flights', askEn: 'Flights from YVR to Phuket in December', labelZh: '普吉岛机票', askZh: '温哥华到普吉岛十二月机票' },
  { emoji: '🇬🇷', labelEn: 'Greece Guide', askEn: 'Best Greek islands to visit? Search Reddit', labelZh: '希腊旅游攻略', askZh: '希腊哪些岛屿最好玩' },
  { emoji: '🇦🇺', labelEn: 'Sydney Flights', askEn: 'Flights from YVR to Sydney Australia', labelZh: '悉尼机票', askZh: '温哥华到悉尼的航班' },
  { emoji: '🇵🇹', labelEn: 'Portugal Guide', askEn: 'What to see in Portugal? WikiVoyage guide', labelZh: '葡萄牙攻略', askZh: '葡萄牙有什么好玩的地方' },
  { emoji: '🇨🇷', labelEn: 'Costa Rica Flights', askEn: 'Flights from YVR to Costa Rica', labelZh: '哥斯达黎加机票', askZh: '温哥华到哥斯达黎加的航班' },
  { emoji: '🇮🇸', labelEn: 'Iceland Guide', askEn: 'Best time to visit Iceland? Trip recommendations', labelZh: '冰岛旅游攻略', askZh: '冰岛最佳旅游时间' },
  { emoji: '🇭🇷', labelEn: 'Croatia Guide', askEn: 'Tell me about visiting Croatia - coast and cities', labelZh: '克罗地亚攻略', askZh: '克罗地亚海岸和城市旅游介绍' },
  { emoji: '🇹🇷', labelEn: 'Turkey Guide', askEn: 'What to see in Istanbul and Turkey? WikiVoyage', labelZh: '土耳其攻略', askZh: '伊斯坦布尔和土耳其有哪些景点' },
  { emoji: '🇯🇵', labelEn: 'Osaka Guide', askEn: 'Tell me about visiting Osaka Japan', labelZh: '大阪旅游', askZh: '日本大阪旅游介绍' },
  { emoji: '🇪🇸', labelEn: 'Barcelona Guide', askEn: 'What to do in Barcelona Spain?', labelZh: '巴塞罗那攻略', askZh: '西班牙巴塞罗那有什么好玩的' },
];

function shufflePrompts() {
  try {
    const isZh = currentLang === 'zh';
    const pool = [...DYN_PROMPTS];
    for (let i = pool.length - 1; i > 0; i--) {
      const j = Math.floor(Math.random() * (i + 1));
      [pool[i], pool[j]] = [pool[j], pool[i]];
    }
    for (let i = 1; i <= 4; i++) {
      const p = pool[i - 1];
      const btn = document.getElementById('dyn' + i);
      if (!btn) continue;
      btn.textContent = p.emoji + ' ' + (isZh ? p.labelZh : p.labelEn);
      btn.onclick = function() { ask(isZh ? p.askZh : p.askEn); };
    }
    // Also translate static presets
    document.querySelectorAll('#presets button:not([id^="dyn"])').forEach(el => {
      if (el.textContent.includes('YVR Deals')) el.textContent = isZh ? '🎫 温哥华特价' : '🎫 YVR Deals';
      else if (el.textContent.includes('Recommend')) el.textContent = isZh ? '🧳 推荐行程' : '🧳 Recommend a Trip';
      else if (el.textContent.includes('Reddit')) el.textContent = isZh ? '🔍 Reddit旅行' : '🔍 Reddit Trips';
      else if (el.textContent.includes('Cruise')) el.textContent = isZh ? '🚢 游轮优惠' : '🚢 Cruise Deals';
    });
  } catch(e) { console.error('shufflePrompts:', e); }
}

applyLang();
refreshSessions();
shufflePrompts();
</script>
</body>
</html>"""

HTML_PAGE = HTML_PAGE.replace("__DEFAULT_MODEL__", DEFAULT_MODEL).replace(
    "__DEFAULT_THINKING__", DEFAULT_THINKING)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.environ.get("PORT", "8092")))
